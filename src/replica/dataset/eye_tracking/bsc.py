"""Beijing Sentence Corpus (Pan et al., 2022).

Chinese sentence-reading eye-tracking corpus published at
https://osf.io/vr3k8/. Two source files are required:

  * ``BSC.Word.Info.v2.xlsx`` -- per-word stimulus table. The sheet
    ``"word"`` has columns ``SN, NW, WORD, POS, LEN, STRK, WF_BLI,
    WF_SUB, PRED, ...``. ``SN`` indexes the sentence (1-150) and
    ``NW`` the within-sentence word position.
  * ``BSC.EMD.txt`` -- per-fixation eye-movement data, tab-separated,
    columns ``id, sn, nw, wn, fl, dur, ao, dir, ...``. Each row is one
    fixation; subject id ``id``, sentence ``sn``, word position ``wn``
    (1-indexed), total word count of the sentence ``nw`` (constant per
    trial), and fixation duration ``dur`` (ms). Rows within a trial are
    in chronological order. (``wn`` is the 1-indexed landing word position;
    ``nw`` the sentence's word count.)

The loader builds one ``Stimulus`` per sentence (text = concatenation of
the Chinese characters; ``words`` annotation layer pinned to the original
``WORD`` segmentation), then for each ``(participant, sentence)`` it
aggregates per-fixation rows into four word-level measures:

  * ``first_fixation``, duration of the first fixation on each word.
  * ``first_run``, the fixations on a word summed from first arrival until
    the eye leaves it; a later return does not contribute.
  * ``first_pass``, whether that arrival fell in first pass; a word first
    entered by a regression is not, so ``first_run_in_first_pass`` treats it as
    a skip.
  * ``total_time``, sum of all fixation durations on each word.

Words that received no fixation are skips, stored as zero.

Reference: Pan et al. (2022). *Behavior Research Methods.*
https://doi.org/10.3758/s13428-021-01730-2
"""

import csv
from collections import defaultdict
from pathlib import Path

import numpy as np

from replica.dataset.base import EyeTrackingDataset
from replica.dataset.contracts import DatasetInfo, Fixation
from replica.dataset.metadata import io as metadata_io

from ...core import AnnotationLayer, Datum, Span, Stimulus
from ..loading.fixation_utils import (
    RunIdInferrer,
    SyntheticTimer,
    append_fixation,
    first_pass_flags,
    first_run_durations,
    mark_unfixated,
)
from ..loading.utils import append_measure_datums

WORD_INFO_FILE = "BSC.Word.Info.v2.xlsx"
EMD_FILE = "BSC.EMD.txt"

_SentenceWords = dict[int, list[tuple[int, str]]]
_StimulusBySentence = dict[int, Stimulus]
_WordIndexBySentence = dict[int, dict[int, int]]
_FixationRows = dict[tuple[str, int], list[tuple[int, int, float]]]
_RawFixationRow = tuple[str, int, int, int, float]


class BSC(EyeTrackingDataset):
    """Beijing Sentence Corpus (Pan et al., 2022).

    Source files (download with ``replica fetch bsc``):

      * ``BSC.Word.Info.v2.xlsx``, sentence and word table.
      * ``BSC.EMD.txt``, per-fixation eye-movement data.

    Default measure: ``first_run_in_first_pass``.
    """

    info = DatasetInfo(name="bsc", language="zh")
    _declared_measures = ("first_fixation", "first_run", "first_pass", "total_time")

    def load(self, path: str) -> None:
        """Load BSC data from a directory containing both source files.

        Args:
            path: Directory holding ``BSC.Word.Info.v2.xlsx`` and
                ``BSC.EMD.txt``.
        """
        self._load_metadata(path, metadata_path=metadata_io.find_dataset_yaml(path))
        word_info, emd = self._source_paths(path)

        sentences = self._read_word_info(word_info)
        stim_by_sn, word_index_by_sn = self._build_stimuli(sentences)
        fixations, raw_rows = self._read_emd_fixations(emd, word_index_by_sn)
        self._append_data(fixations, stim_by_sn, word_index_by_sn)
        self._append_raw_fixations(raw_rows, stim_by_sn, word_index_by_sn)

    @staticmethod
    def _source_paths(path: str) -> tuple[Path, Path]:
        root = Path(path)
        word_info = root / WORD_INFO_FILE
        emd = root / EMD_FILE
        if not word_info.exists():
            raise FileNotFoundError(
                f"BSC: {WORD_INFO_FILE} not found in {root}. Run `replica fetch bsc`."
            )
        if not emd.exists():
            raise FileNotFoundError(
                f"BSC: {EMD_FILE} not found in {root}. Run `replica fetch bsc`."
            )
        return word_info, emd

    def _build_stimuli(
        self, sentences: _SentenceWords
    ) -> tuple[_StimulusBySentence, _WordIndexBySentence]:
        self._stimuli = []
        stim_by_sn: _StimulusBySentence = {}
        word_index_by_sn: _WordIndexBySentence = {}
        for sn in sorted(sentences):
            words = sentences[sn]  # list[(nw, word_str)] sorted by nw
            text = "".join(w for _, w in words)
            spans, idx_map, offset = [], {}, 0
            for i, (nw, w) in enumerate(words):
                start = offset
                end = offset + len(w)
                spans.append(Span(start, end))
                idx_map[nw] = i
                offset = end
            layer = AnnotationLayer(
                name="words", spans=spans, values=[w for _, w in words]
            )
            stim = Stimulus(
                text=text,
                stimulus_id=f"bsc_s{sn}",
                metadata={"sentence_num": sn},
            )
            stim.add_layer(layer)
            self._stimuli.append(stim)
            stim_by_sn[sn] = stim
            word_index_by_sn[sn] = idx_map
        return stim_by_sn, word_index_by_sn

    @staticmethod
    def _read_emd_fixations(
        emd: Path, word_index_by_sn: _WordIndexBySentence
    ) -> tuple[_FixationRows, list[_RawFixationRow]]:
        # EMD column semantics (verified empirically against Word.Info.v2.xlsx):
        #   ``wn`` = 1-indexed word position the fixation landed on
        #   ``nw`` = total word count of the sentence (constant per trial),
        #            not read by this parser; word geometry comes from
        #            Word.Info via ``word_index_by_sn``
        # Row order within a trial in the file is chronological and there is no
        # per-fixation index column, so a synthetic ``fix_seq`` counter is
        # assigned while reading to let the later sort preserve that order.
        fixations: _FixationRows = defaultdict(list)
        raw_rows: list[_RawFixationRow] = []
        fix_seq_per_trial: dict[tuple[str, int], int] = defaultdict(int)
        with open(emd, encoding="utf-8") as f:
            reader = csv.DictReader(f, delimiter="\t")
            for row in reader:
                parsed = BSC._parse_emd_row(row)
                if parsed is None:
                    continue
                sid, sn, word_pos, dur = parsed
                if sn not in word_index_by_sn:
                    continue
                if word_pos not in word_index_by_sn[sn]:
                    continue
                key = (sid, sn)
                seq = fix_seq_per_trial[key]
                fix_seq_per_trial[key] = seq + 1
                fixations[key].append((seq, word_pos, dur))
                raw_rows.append((str(sid), sn, word_pos, seq, dur))
        return fixations, raw_rows

    @staticmethod
    def _parse_emd_row(row: dict[str, str]) -> tuple[str, int, int, float] | None:
        try:
            return row["id"], int(row["sn"]), int(row["wn"]), float(row["dur"])
        except (KeyError, ValueError):
            return None

    def _append_data(
        self,
        fixations: _FixationRows,
        stim_by_sn: _StimulusBySentence,
        word_index_by_sn: _WordIndexBySentence,
    ) -> None:
        self._data: list[Datum] = []
        for (sid, sn), fixes in fixations.items():
            stim = stim_by_sn[sn]
            n_words = len(stim.layers["words"])
            idx_map = word_index_by_sn[sn]
            fixes.sort()  # by fix_seq ascending = chronological (file order)

            ff = np.full(n_words, np.nan)  # first_fixation
            tt = np.full(n_words, np.nan)  # total_time
            fr = first_run_durations(
                [word_pos for _seq, word_pos, _dur in fixes],
                [dur for _seq, _word_pos, dur in fixes],
                idx_map,
                n_words,
            )
            fp = first_pass_flags(
                [idx_map[word_pos] for _seq, word_pos, _dur in fixes], n_words
            )

            seen_first = set()
            for _seq, word_pos, dur in fixes:
                idx = idx_map[word_pos]
                tt[idx] = (0.0 if np.isnan(tt[idx]) else tt[idx]) + dur
                if idx not in seen_first:
                    ff[idx] = dur
                    seen_first.add(idx)

            measures = (
                ("first_fixation", ff),
                ("first_run", fr),
                ("first_pass", fp),
                ("total_time", tt),
            )
            mark_unfixated(
                [ff, fr, tt],
                idx_map.values(),
                (idx_map[word_pos] for _, word_pos, _ in fixes),
            )
            append_measure_datums(
                self._data,
                stimulus=stim,
                participant_id=str(sid),
                measure_names=[name for name, _ in measures],
                measure_values=[values for _, values in measures],
            )

    def _append_raw_fixations(
        self,
        raw_rows: list[_RawFixationRow],
        stim_by_sn: _StimulusBySentence,
        word_index_by_sn: _WordIndexBySentence,
    ) -> None:
        # 3. Emit raw fixations. BSC's EMD has ``dur`` but no per-fixation
        # start/end timestamps, so trial-relative timing is synthesised.
        # ``run_id`` is inferred per-IA from the chronological stream.
        # raw_rows tuple layout: (sid, sn, word_pos, fix_seq, dur).
        self._fixations: list[Fixation] = []
        timer = SyntheticTimer()
        run_ids = RunIdInferrer()
        # Sort raw_rows so that within each (sid, sn) trial fixations are
        # processed in chronological order (by the ``fix_seq`` counter
        # assigned while reading the file).
        raw_rows.sort(key=lambda r: (r[0], r[1], r[3]))
        for sid, sn, word_pos, fix_seq, dur in raw_rows:
            stim = stim_by_sn[sn]
            idx = word_index_by_sn[sn][word_pos]
            char_offset = stim.layers["words"].spans[idx].start
            trial_key = (sid, sn)
            start, end = timer.next(trial_key, dur)
            run_id = run_ids.next(trial_key, idx)
            append_fixation(
                self._fixations,
                participant_id=sid,
                stimulus_id=stim.stimulus_id,
                start_time=start,
                end_time=end,
                duration=dur,
                x=float("nan"),
                y=float("nan"),
                unit_index=idx,
                char_offset=char_offset,
                fixation_index=fix_seq + 1,
                metadata={
                    "sentence_num": sn,
                    "word_pos": word_pos,
                    "run_id": run_id,
                    "synthetic_start_time": True,
                },
            )

    def _read_word_info(self, xlsx_path: Path) -> dict[int, list[tuple[int, str]]]:
        """Parse the per-word stimulus table from BSC.Word.Info.v2.xlsx."""
        import openpyxl

        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        if "word" not in wb.sheetnames:
            raise ValueError(
                f"Sheet 'word' missing from {xlsx_path.name}; sheets: {wb.sheetnames}"
            )
        ws = wb["word"]
        rows = ws.iter_rows(values_only=True)
        header = [h for h in next(rows)]
        try:
            sn_col = header.index("SN")
            nw_col = header.index("NW")
            word_col = header.index("WORD")
        except ValueError as e:
            raise ValueError(
                f"BSC word sheet missing expected SN/NW/WORD columns: {header}"
            ) from e

        sentences: dict[int, list[tuple[int, str]]] = defaultdict(list)
        for row in rows:
            if row[sn_col] is None:
                continue
            try:
                sn = int(row[sn_col])
                nw = int(row[nw_col])
            except (TypeError, ValueError):
                continue
            word = row[word_col]
            if word is None:
                continue
            sentences[sn].append((nw, str(word)))
        for sentence_rows in sentences.values():
            sentence_rows.sort()
        return dict(sentences)
